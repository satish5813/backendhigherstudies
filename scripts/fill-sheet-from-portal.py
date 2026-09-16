"""
Fill a placement spreadsheet with what the portal knows about each student.

    python scripts/fill-sheet-from-portal.py --sheet "path/to/batch.xlsx" \
        --export path/to/portal-export.json [--id-column "Student ID NO"] [--out path/to/output.xlsx]

The export JSON comes from GET /api/cohorts/export (see server/src/db/pull-portal-data.js).
Rows are matched on the registration number. Students the portal has never
heard of get empty cells — nothing is guessed. The original file is never
modified; the output is a copy with new columns appended after the last
existing one, styled like the existing header.
"""
import argparse
import json
import re
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Column header -> function(row) -> value. Order is the order in the sheet.
COLUMNS = [
    ("On roster (cohort)",        lambda r: r.get("cohort")),
    ("Campus",                    lambda r: r.get("campus")),
    ("Profile created",           lambda r: "Yes" if r.get("account") else "No"),
    ("Portal login email",        lambda r: r.get("loginEmail")),
    ("Last sign-in",              lambda r: _date(r.get("account", {}) and r["account"].get("lastLoginAt"))),
    ("Onboarding done",           lambda r: _yn(r.get("account", {}) and r["account"].get("onboarded"))),
    ("Public profile",            lambda r: r.get("account", {}) and r["account"].get("publicUrl")),
    ("Headline",                  lambda r: _p(r, "headline")),
    ("Location",                  lambda r: _p(r, "location")),
    ("CGPA (portal)",             lambda r: _p(r, "cgpa")),
    ("Graduation year",           lambda r: _p(r, "gradYear")),
    ("Skills",                    lambda r: ", ".join(_p(r, "skills") or [])),
    ("GitHub",                    lambda r: _link(r, "github")),
    ("LinkedIn",                  lambda r: _link(r, "linkedin")),
    ("LeetCode",                  lambda r: _link(r, "leetcode")),
    ("Portfolio",                 lambda r: _link(r, "portfolio")),
    ("LeetCode solved",           lambda r: _coding(r, "leetcode", "solved")),
    ("LeetCode rating",           lambda r: _coding(r, "leetcode", "rating")),
    ("CodeChef rating",           lambda r: _coding(r, "codechef", "rating")),
    ("Projects",                  lambda r: _p(r, "projects")),
    ("Project titles",            lambda r: "; ".join(_p(r, "projectTitles") or [])),
    ("Internships / experience",  lambda r: _p(r, "experiences")),
    ("Achievements",              lambda r: _p(r, "achievements")),
    ("Resumes",                   lambda r: (r.get("resume") or {}).get("count")),
    ("Best resume",               lambda r: (r.get("resume") or {}).get("title")),
    ("ATS score",                 lambda r: (r.get("resume") or {}).get("atsScore")),
    ("Resume updated",            lambda r: _date((r.get("resume") or {}).get("updatedAt"))),
    ("Applications recorded",     lambda r: r.get("applications")),
    ("Target band (placement cell)", lambda r: r.get("band")),
    ("Readiness index",           lambda r: r.get("readiness")),
]


def _p(r, key):
    return (r.get("profile") or {}).get(key)


def _yn(v):
    if v is None or v == "":
        return None
    return "Yes" if v else "No"


def _date(v):
    if not v:
        return None
    try:
        return datetime.fromisoformat(str(v).replace("Z", "+00:00")).strftime("%d %b %Y")
    except ValueError:
        return str(v)


def _link(r, platform):
    for l in _p(r, "links") or []:
        if l.get("platform") == platform:
            return l.get("url") or l.get("username")
    return None


def _coding(r, platform, field):
    for c in _p(r, "coding") or []:
        if c.get("platform") == platform:
            return c.get(field)
    return None


def norm(v):
    return re.sub(r"[^A-Za-z0-9]", "", str(v or "")).upper()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--export", required=True)
    ap.add_argument("--id-column", default="Student ID NO")
    ap.add_argument("--out")
    ap.add_argument("--worksheet", help="sheet name; default: the first")
    a = ap.parse_args()

    data = json.loads(Path(a.export).read_text(encoding="utf-8"))
    by_reg = {norm(row["regNo"]): row for row in data["items"]}

    wb = openpyxl.load_workbook(a.sheet)
    ws = wb[a.worksheet] if a.worksheet else wb.worksheets[0]

    headers = [c.value for c in ws[1]]
    try:
        id_col = next(i for i, h in enumerate(headers, start=1) if h and str(h).strip().lower() == a.id_column.strip().lower())
    except StopIteration:
        sys.exit(f"No column named {a.id_column!r} in row 1. Found: {headers}")

    # Style new headers like the existing last header cell.
    last = ws.max_column
    template = ws.cell(row=1, column=last)
    start = last + 1
    for offset, (title, _) in enumerate(COLUMNS):
        cell = ws.cell(row=1, column=start + offset, value=title)
        cell.font = copy(template.font)
        cell.fill = copy(template.fill)
        cell.border = copy(template.border)
        cell.alignment = copy(template.alignment)
        ws.column_dimensions[get_column_letter(start + offset)].width = max(14, min(42, len(title) + 4))

    body_font = copy(ws.cell(row=2, column=last).font) if ws.max_row >= 2 else Font(name="Arial", size=10)

    matched = 0
    for row in range(2, ws.max_row + 1):
        reg = ws.cell(row=row, column=id_col).value
        if reg is None or str(reg).strip() == "":
            continue
        portal = by_reg.get(norm(reg))
        if portal:
            matched += 1
        for offset, (_, fn) in enumerate(COLUMNS):
            value = fn(portal) if portal else None
            if value in ("", [], {}):
                value = None
            cell = ws.cell(row=row, column=start + offset, value=value)
            cell.font = copy(body_font)
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 30)

    # A note at the end of the table saying where the new columns came from,
    # so the sheet explains itself to the next person who opens it.
    note_row = ws.max_row + 2
    ws.cell(row=note_row, column=1, value=(
        f"Columns from '{COLUMNS[0][0]}' onward were filled from the KL Placement Readiness portal on "
        f"{datetime.now().strftime('%d %b %Y %H:%M')} by matching '{a.id_column}' to the portal's registration number. "
        f"Blank = the portal has no such student, or the student has not entered that detail yet."
    )).font = Font(name=body_font.name or "Arial", size=9, italic=True, color="666666")

    out = a.out or str(Path(a.sheet).with_name(Path(a.sheet).stem + " - portal data.xlsx"))
    wb.save(out)
    total = sum(1 for r in range(2, ws.max_row) if ws.cell(row=r, column=id_col).value not in (None, ""))
    print(f"matched {matched} of {total} students; wrote {out}")


if __name__ == "__main__":
    main()
